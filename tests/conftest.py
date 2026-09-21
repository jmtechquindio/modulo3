"""Fixtures compartidos de pruebas.

Se construye un XLSM de prueba con la misma estructura que el archivo real
(encabezados en fila 4, datos desde fila 5) y se migra a una base temporal.
Permite probar el importador y las herramientas sin depender del archivo original.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from scripts import importar_excel  # noqa: E402

ENC_PROVEEDORES = importar_excel.ENCABEZADOS["Proveedores"]
ENC_PRODUCTOS = importar_excel.ENCABEZADOS["Productos"]
ENC_INVENTARIO = importar_excel.ENCABEZADOS["Inventario"]

FECHA_POR_VENCER = "2026-09-10"  # a 6 días del corte 2026-09-04
FECHA_VENCIDO = "2025-04-28"


def construye_hoja(ws, titulo: str, encabezados: list[str], datos: list[tuple],
                   control: str | None = None, valor_control=None):
    fila2 = "Datos ficticios | prueba"
    if control:
        fila2 += f" | {control}"
    ws.append([titulo])
    ws.append([fila2])
    fila3 = [""] * len(encabezados)
    if control:
        fila3[-2] = control
        fila3[-1] = valor_control
    ws.append(fila3)
    ws.append(list(encabezados))
    for fila in datos:
        ws.append(list(fila))


def construir_xlsm(destino: Path, *, proveedor_invalido: bool = False,
                   columna_ausente: bool = False) -> Path:
    from datetime import date

    wb = openpyxl.Workbook()

    proveedores = wb.active
    proveedores.title = "Proveedores"
    construye_hoja(
        proveedores, "Maestro de proveedores", ENC_PROVEEDORES,
        [("PRV-001", "Distribuciones Andina FER S.A.S.", "900000000-3", "Ferretería",
          "Ricardo Rodríguez", "6036901859", "3015085640",
          "ricardo@distribucionesandina.co", "Carrera 42 # 52-54", "Bogotá D.C.",
          "Cundinamarca", "Contado", "Transferencia bancaria", "Activo")],
        control="Total proveedores", valor_control=1,
    )

    productos = wb.create_sheet("Productos")
    datos_productos = [
        ("FER-0001", "Martillo de uña Forte Mini", "Ferretería", "PRV-001",
         "Distribuciones Andina FER S.A.S.", "Unidad", 98600, 121300, 0.19, None),
        ("ABA-0001", "Caldo concentrado Rico Día Premium", "Abarrotes", "PRV-001",
         "Distribuciones Andina FER S.A.S.", "Unidad", 32700, 40000, 0.22, None),
    ]
    if proveedor_invalido:
        datos_productos.append(
            ("XZZ-0001", "Producto con proveedor inválido", "Otra", "PRV-999",
             "Proveedor Fantasma S.A.S.", "Unidad", 100, 150, 0.5, None)
        )
    prod_headers = list(ENC_PRODUCTOS)
    if columna_ausente:
        prod_headers.remove("Precio costo")
    construye_hoja(productos, "Catálogo maestro de productos", prod_headers,
                   datos_productos, control="Total productos",
                   valor_control=len(datos_productos))

    inventario = wb.create_sheet("Inventario")
    inv_headers = list(ENC_INVENTARIO)
    if columna_ausente:
        inv_headers.remove("Stock mínimo")
    construye_hoja(
        inventario, "Inventario disponible", inv_headers,
        [
            ("FER-0001", "Martillo de uña Forte Mini", "Ferretería",
             "Distribuciones Andina FER S.A.S.", "D-14-06", "LT-2430339", 8, 22,
             98600, 788800, date.fromisoformat(FECHA_POR_VENCER), None, "POR VENCER",
             date.fromisoformat("2026-08-07"), "Priorizar rotación"),
            ("ABA-0001", "Caldo concentrado Rico Día Premium", "Abarrotes",
             "Distribuciones Andina FER S.A.S.", "A-02-01", "LT-2602617", 50, 10,
             32700, 1635000, None, None, "OK", None, None),
            ("ABA-0001", "Caldo concentrado Rico Día Premium", "Abarrotes",
             "Distribuciones Andina FER S.A.S.", "A-02-02", "LT-VENCIDO", 12, 10,
             32700, 392400, date.fromisoformat(FECHA_VENCIDO), None, "VENCIDO",
             None, "Separar"),
        ],
        control="Productos inventariados", valor_control=3,
    )
    wb.save(destino)
    return destino


@pytest.fixture
def base(tmp_path: Path):
    db = tmp_path / "inventario.db"
    importar_excel.construir_base(construir_xlsm(tmp_path / "prueba.xlsm"), db)
    return db


@pytest.fixture
def base_con_proveedor_invalido(tmp_path: Path):
    db = tmp_path / "inventario_invalido.db"
    return importar_excel.construir_base(
        construir_xlsm(tmp_path / "prueba_invalida.xlsm", proveedor_invalido=True), db
    )